"""Parse the Texas refinery fenceline benzene reports (CEDRI 63.655(h)(8) template) for check V3.

Each zip in data/external/epa_fenceline/reports holds one or more xlsx workbooks (some hold nested zips). From each
workbook it reads, by the template's machine keys rather than cell positions:
  Facility Information:   FacilityName, StateFacID (usually the TCEQ RN), SamplingYear, SamplingQuarter
  Period Identifications: PeriodStartDate, PeriodEndDate, PeriodAvg (biweekly delta c), AnnualAvg (rolling delta c)
  Sample Results:         number of sampler rows and of rows flagged below detection, per period
Repeated submissions of the same period (resubmissions, one file per attachment) are collapsed: for each facility and
period the value from the latest submission is kept, and disagreements between submissions are counted.

Outputs: data/derived/fenceline_periods.csv, data/derived/fenceline_parse_log.json
Usage: python3 src/parse_fenceline.py
"""
import io
import json
import re
import warnings
import zipfile
from pathlib import Path

import openpyxl
import pandas as pd

warnings.filterwarnings("ignore")
ROOT = Path(__file__).resolve().parent.parent
REP = ROOT / "data" / "external" / "epa_fenceline" / "reports"
DER = ROOT / "data" / "derived"


def sheet_records(ws, key):
    """Rows under the machine-key header row (the row holding `key`), skipping the 'e.g.:' example rows."""
    rows = list(ws.iter_rows(values_only=True))
    for i, r in enumerate(rows):
        if r and key in [str(c).strip() if c is not None else None for c in r]:
            hdr = [str(c).strip() if c is not None else None for c in r]
            out = []
            for rr in rows[i + 1:]:
                vals = dict(zip(hdr, rr))
                cells = [v for k, v in vals.items() if k and v not in (None, "")]
                if not cells or any(str(v).startswith("e.g.") for v in cells):
                    continue
                out.append({k: v for k, v in vals.items() if k})
            return out
    return None


def workbooks(path):
    def walk(z, label):
        for n in z.namelist():
            low = n.lower()
            if low.endswith(".xlsx"):
                yield f"{label}/{n}", z.read(n)
            elif low.endswith(".zip"):
                yield from walk(zipfile.ZipFile(io.BytesIO(z.read(n))), f"{label}/{n}")
    yield from walk(zipfile.ZipFile(path), path.name)


def num(v):
    try:
        return float(str(v).replace(",", "").strip())
    except ValueError:
        return None


def main():
    idx = pd.read_csv(ROOT / "data" / "external" / "epa_fenceline" / "reports_index.csv")
    sub = dict(zip(idx.file.str.split("/").str[-1], idx.submitted))
    rows, log = [], {"zips": 0, "workbooks": 0, "not_template": [], "no_periods": []}
    for z in sorted(REP.glob("*/*.zip")):
        log["zips"] += 1
        for label, data in workbooks(z):
            log["workbooks"] += 1
            try:
                wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
            except Exception:
                log["not_template"].append(label)
                continue
            names = wb.sheetnames
            if "Period Identifications" not in names or "Facility Information" not in names:
                log["not_template"].append(label)
                continue
            fac = sheet_records(wb["Facility Information"], "FacilityName") or [{}]
            per = sheet_records(wb["Period Identifications"], "PeriodAvg") or []
            res = sheet_records(wb["Sample Results"], "BenzeneAmt") if "Sample Results" in names else None
            if not per:
                log["no_periods"].append(label)
            nres = pd.DataFrame(res or [])
            for p in per:
                pid = str(p.get("PeriodId", "")).strip()
                r = nres[nres.get("PeriodId", pd.Series(dtype=str)).astype(str).str.strip() == pid] if len(nres) else nres
                rows.append({
                    "zip": z.name, "workbook": label, "submitted": sub.get(z.name),
                    "facility": str(fac[0].get("FacilityName", "")).strip(),
                    "state_fac_id": str(fac[0].get("StateFacID", "")).strip().upper(),
                    "period_id": pid, "start": p.get("PeriodStartDate"), "end": p.get("PeriodEndDate"),
                    "period_dc": num(p.get("PeriodAvg")), "annual_dc": num(p.get("AnnualAvg")),
                    "period_comment": p.get("PeriodComments"),
                    "n_samples": len(r), "n_bdl": int(r.get("BdlFlag", pd.Series(dtype=str)).astype(str).str.lower()
                                                     .str.startswith("y").sum()) if len(r) else 0})
    d = pd.DataFrame(rows)
    d["start"] = pd.to_datetime(d.start, errors="coerce", format="mixed")
    d["end"] = pd.to_datetime(d.end, errors="coerce", format="mixed")
    d["rn"] = d.state_fac_id.str.extract(r"(RN\d{9})", expand=False)
    # facility identity from the hand-made crosswalk (data/manual/fenceline_crosswalk.csv, each row with its basis)
    cw = pd.read_csv(ROOT / "data" / "manual" / "fenceline_crosswalk.csv")
    d = d.merge(cw.rename(columns={"rn": "rn_crosswalk"})[["workbook_facility", "refinery", "network", "rn_crosswalk"]],
                left_on="facility", right_on="workbook_facility", how="left")
    # the Borger filings carry two sampler rings; the JTF ring's zips say JTF but some of its workbooks reuse the
    # main ring's facility name, so the zip name decides
    d.loc[d.refinery.eq("borger") & d.zip.str.contains("JTF", case=False), "network"] = "jtf"
    log["rows_without_crosswalk"] = int(d.workbook_facility.isna().sum())
    d = d[d.refinery.notna()].sort_values("submitted")
    key = ["refinery", "network", "start", "end"]
    disagree = d.dropna(subset=["period_dc"]).groupby(key).period_dc.nunique()
    out = d.drop_duplicates(key, keep="last")
    days = (out.end - out.start).dt.days
    out = out.assign(period_days=days, date_suspect=~days.between(10, 18))
    out.to_csv(DER / "fenceline_periods.csv", index=False, date_format="%Y-%m-%d")
    log.update({"period_rows_read": len(d), "unique_facility_periods": len(out),
                "facility_periods_with_disagreeing_values": int((disagree > 1).sum()),
                "refineries": int(out.refinery.nunique()), "sampler_networks": int(out.groupby(["refinery", "network"]).ngroups),
                "periods_per_network": out.groupby(["refinery", "network"]).size().describe().round(1).to_dict(),
                "rn_stated_differs_from_crosswalk": int((out.rn.notna() & (out.rn != out.rn_crosswalk)).sum()),
                "first_period": str(out.start.min()), "last_period": str(out.end.max()),
                "unparsed_dates": int(out.start.isna().sum()), "date_suspect": int(out.date_suspect.sum()), "missing_period_dc": int(out.period_dc.isna().sum()),
                "not_template_count": len(log["not_template"]), "no_periods_count": len(log["no_periods"])})
    (DER / "fenceline_parse_log.json").write_text(json.dumps(log, indent=1, default=str) + "\n")
    print(json.dumps({k: v for k, v in log.items() if not isinstance(v, list)}, indent=1, default=str))


if __name__ == "__main__":
    main()
